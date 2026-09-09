from PyQt5.QtWidgets import QWidget, QPushButton, QFileDialog, QTableWidget, QLabel, QLineEdit

from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as letter
from openpyxl import Workbook

from src.variables import *

import hashlib
import json


class TabExport(QWidget):
    def __init__(self, window):
        super().__init__()

        self.window = window

        self.objects = {}
        self.classes = []

        for i, cnt in enumerate(self.window.settings["classes"]["count"]):
            for number in range(cnt):
                self.classes.append(f"{i + 1} {CLASSES_ALPHABET[number + 1]}")

        self.timesRama = QTableWidget(0, 0, parent=self)
        self.timesRama.horizontalHeader().setVisible(False)
        self.timesRama.verticalHeader().setVisible(False)
        self.timesRama.show()

        lessons = self.window.settings["max_lesson_count_per_day"]
        shifts = self.window.settings["number_of_shifts"]

        flag = False

        for shift in range(shifts):
            for lesson in range(lessons):
                if f"{shift}-{lesson}" not in self.window.settings["display"]["time"]:
                    self.window.settings["display"]["time"][f"{shift}-{lesson}"] = "00:00 - 00:00"

                    flag = True

                self.objects[f"time_{shift}-{lesson}_label"] = QLabel(parent=self)
                self.objects[f"time_{shift}-{lesson}_label"].setFont(FONT)
                self.objects[f"time_{shift}-{lesson}_label"].setText(f"{translate('menu.main.tab.export.shift')} - {shift + 1}, {translate('menu.main.tab.export.lesson')} - {lesson + 1}:")
                self.objects[f"time_{shift}-{lesson}_label"].show()

                self.objects[f"time_{shift}-{lesson}_lineedit"] = QLineEdit(parent=self)
                self.objects[f"time_{shift}-{lesson}_lineedit"].setFont(FONT)
                self.objects[f"time_{shift}-{lesson}_lineedit"].setText(self.window.settings["display"]["time"][f"{shift}-{lesson}"])
                self.objects[f"time_{shift}-{lesson}_lineedit"].editingFinished.connect(lambda s = shift, l = lesson: self.timeLineEditEditingFinished(s, l))
                self.objects[f"time_{shift}-{lesson}_lineedit"].show()

        if flag:
            with open(f"{PATH_TO_FOLDER}/projects/{self.window.project}/settings.json", "w", encoding="utf-8") as file:
                json.dump(self.window.settings, file, ensure_ascii=False)

        self.exportByClassPushButton = QPushButton(parent=self)
        self.exportByClassPushButton.setFont(FONT)
        self.exportByClassPushButton.clicked.connect(lambda: self.exportByClassPushButtonClicked())
        self.exportByClassPushButton.setText(translate("menu.main.tab.export.save_classes_schedule"))
        self.exportByClassPushButton.show()

        self.exportByTeacherPushButton = QPushButton(parent=self)
        self.exportByTeacherPushButton.setFont(FONT)
        self.exportByTeacherPushButton.clicked.connect(lambda: self.exportByTeacherPushButtonClicked())
        self.exportByTeacherPushButton.setText(translate("menu.main.tab.export.save_teachers_schedule"))
        self.exportByTeacherPushButton.show()

    def timeLineEditEditingFinished(self, shift, lesson):
        self.window.settings["display"]["time"][f"{shift}-{lesson}"] = self.objects[f"time_{shift}-{lesson}_lineedit"].text()

        with open(f"{PATH_TO_FOLDER}/projects/{self.window.project}/settings.json", "w", encoding="utf-8") as file:
            json.dump(self.window.settings, file, ensure_ascii=False)

    def start(self, page):
        page["A1"] = translate("menu.main.tab.export.day")
        page["B1"] = "#"
        page["C1"] = translate("menu.main.tab.export.time")

        lessons = self.window.settings["max_lesson_count_per_day"]
        shifts = self.window.settings["number_of_shifts"]

        height = lessons * shifts

        lines = []

        for day in range(self.window.settings["working_days_per_week"]):
            lesson = 0

            for lesson in range(height):
                page[f"B{day * (height + 1) + lesson + 2}"] = lesson % lessons + 1

            page.merge_cells(f"A{day * (height + 1) + 2}:A{day * (height + 1) + lesson + 2}")
            page[f"A{day * (height + 1) + 2}"] = translate(f"abbreviate.day.{day}")

            lines.append(day * (height + 1) + lesson + 3)

            for shift in range(shifts):
                for lesson in range(lessons):
                    page[f"C{day * (height + 1) + shift * lessons + lesson + 2}"] = self.window.settings["display"]["time"][f"{shift}-{lesson}"]

        return height, lines

    def finish(self, page, lines):
        page.row_dimensions[1].height = 20

        for col in page.columns:
            mx = 0

            for cell in col:
                cell.alignment = Alignment(horizontal="center", vertical="center")

                if cell.value is not None and len(str(cell.value)) > mx:
                    mx = len(str(cell.value))

            page.column_dimensions[letter(col[0].column)].width = mx + 2

        for line in lines:
            page.row_dimensions[line].height = 3

            for col in range(1, page.max_column + 1):
                page.cell(row=line, column=col).fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    @staticmethod
    def color(text: str) -> str:
        number = (int.from_bytes(hashlib.sha256(text.encode()).digest(), byteorder="big") % 16777216) % 0x1000000
        hexnum = f"{number:06x}"

        return "".join(f"{int((int(hexnum[i:i + 2], 16) + 255) / 2 + 0.5):02x}" for i in range(0, 6, 2))

    def exportByClassPushButtonClicked(self, path: str = None):
        if path is None:
            path, _ = QFileDialog.getSaveFileName(self, translate("menu.main.tab.export.save_file"), "classes.xlsx", "Excel (*.xlsx)")

        if not path:
            return

        if not os.path.exists(f"{PATH_TO_FOLDER}/projects/{self.window.project}/answer.json"):
            return

        with open(f"{PATH_TO_FOLDER}/projects/{self.window.project}/answer.json", "r", encoding="utf-8") as file:
            answer = json.load(file)

        book = Workbook()

        page = book.active
        page.title = translate("menu.main.tab.export.schedule")

        height, lines = self.start(page)

        lessons = self.window.settings["max_lesson_count_per_day"]

        for col, cls in enumerate(self.classes, start=4):
            page[f"{letter(col)}1"] = cls

        for cls, days in answer.items():
            pos = self.classes.index(cls) + 4

            shift = self.window.settings["classes"]["shift"][int(cls.split(" ")[0]) - 1]

            for day in range(len(days)):
                for i, lesson in enumerate(days[day]):
                    if lesson["subject"] == "#":
                        continue

                    subjects = ", ".join([lesson["subject"]] + [element["subject"] for element in lesson["extra"]])

                    col = pos
                    row = day * (height + 1) + shift * lessons + i + 2

                    page[f"{letter(col)}{row}"] = subjects

                    color = self.color(subjects)

                    page.cell(row=row, column=col).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        self.finish(page, lines)

        book.save(path)

    def exportByTeacherPushButtonClicked(self, path: str = None):
        if path is None:
            path, _ = QFileDialog.getSaveFileName(self, translate("menu.main.tab.export.save_file"), "teachers.xlsx", "Excel (*.xlsx)")

        if not path:
            return

        if not os.path.exists(f"{PATH_TO_FOLDER}/projects/{self.window.project}/answer.json"):
            return

        with open(f"{PATH_TO_FOLDER}/projects/{self.window.project}/answer.json", "r", encoding="utf-8") as file:
            answer = json.load(file)

        teachers = set()

        for cls, days in answer.items():
            for day in days:
                for lesson in day:
                    for element in [lesson] + lesson["extra"]:
                        if element["subject"] == "#":
                            continue

                        teachers.update(element.get("teachers", []))

        teachers = sorted(teachers)

        book = Workbook()

        page = book.active
        page.title = translate("menu.main.tab.export.schedule")

        height, lines = self.start(page)

        lessons = self.window.settings["max_lesson_count_per_day"]

        for col, teacher in enumerate(teachers, start=4):
            page[f"{letter(col)}1"] = teacher

        for cls, days in answer.items():
            shift = self.window.settings["classes"]["shift"][int(cls.split(" ")[0]) - 1]

            for day in range(len(days)):
                for i, lesson in enumerate(days[day]):
                    row = day * (height + 1) + shift * lessons + i + 2

                    for element in [lesson] + lesson["extra"]:
                        if element["subject"] == "#" or not element.get("teachers"):
                            continue

                        content = f"{cls}: {element['subject']}"
                        color = self.color(element["subject"])

                        for teacher in element["teachers"]:
                            col = teachers.index(teacher) + 4

                            cell = page.cell(row=row, column=col)

                            if cell.value:
                                cell.value = f"{cell.value}; {content}"

                            else:
                                cell.value = content

                            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        self.finish(page, lines)

        book.save(path)
